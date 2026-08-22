"""
genera_lista_mayorista.py — Stockroom

Lista de precios MAYORISTA para pasarle a un cliente.

Uso: python genera_lista_mayorista.py --output lista.xlsx
     [--ratio 0.773] [--desc 50:5 100:10] [--entrega-dias 8]
     [--validez 15] [--incluir-protectores]

IMPORTANTE: este archivo SE LE ENTREGA A UN TERCERO. No lleva costos, márgenes,
precios de Mercado Libre, proveedores ni existencias. Solo modelo, foto,
colores, compatibilidad y precio mayorista.

Cada modelo tiene su PROPIO precio (precio público × ratio); el descuento por
volumen es una FÓRMULA sobre el total de unidades, así que el cliente ve bajar
el precio de cada fila a medida que carga cantidades.
"""
import argparse, io, json, os, re, sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def _norm(s):
    s = str(s or '').lower()
    for a, b in zip('áéíóúüñ', 'aeiouun'):
        s = s.replace(a, b)
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def _miniatura(url, cache, lado=96):
    """PNG cuadrado listo para incrustar. Cachea BYTES: openpyxl lee las
    imágenes recién al guardar y PIL cierra el buffer, así que dos celdas no
    pueden compartir el mismo BytesIO."""
    if not url:
        return None
    if url in cache:
        d = cache[url]
        return io.BytesIO(d) if d is not None else None
    try:
        import urllib.request
        from PIL import Image as PILImage
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=8) as r:
            raw = r.read()
        im = PILImage.open(io.BytesIO(raw)).convert('RGB')
        im.thumbnail((lado, lado), PILImage.LANCZOS)
        lienzo = PILImage.new('RGB', (lado, lado), (255, 255, 255))
        lienzo.paste(im, ((lado - im.width) // 2, (lado - im.height) // 2))
        buf = io.BytesIO()
        lienzo.save(buf, format='PNG')
        cache[url] = buf.getvalue()
        return io.BytesIO(cache[url])
    except Exception as e:
        print(f"  [foto] {url[:60]}: {e}", file=sys.stderr)
        cache[url] = None
        return None


def leer_modelos(base, incluir_protectores=False):
    """Modelos únicos de funda con stock, deduplicados entre las dos cuentas.

    La misma publicación vive en las dos cuentas (stock sincronizado), así que
    contar los dos caches sin deduplicar duplicaría el catálogo entero.
    """
    from lib_fundas import es_funda
    modelos, vistas = {}, set()
    cache_dir = os.path.join(base, 'cache')
    for f in sorted(os.listdir(cache_dir)):
        if not re.match(r'^items-.+\.json$', f):
            continue
        try:
            raw = json.load(open(os.path.join(cache_dir, f), encoding='utf-8'))
        except Exception:
            continue
        items = raw if isinstance(raw, list) else raw.get('items', [])
        for it in items:
            if not it or not it.get('id'):
                continue
            titulo = it.get('title') or ''
            if not es_funda(titulo):
                continue
            if not incluir_protectores and 'protector' in titulo.lower():
                continue
            if str(it.get('status') or 'active') != 'active':
                continue
            k = _norm(titulo)
            m = modelos.setdefault(k, {
                'titulo': titulo, 'colores': [], 'equipos': [],
                'stock': 0, 'foto': '', 'precio_ml': 0,
            })
            if (it.get('price') or 0) > m['precio_ml']:
                m['precio_ml'] = it.get('price') or 0
            if not m['foto']:
                pics = it.get('pictures') or []
                m['foto'] = (pics[0].get('secure_url') or pics[0].get('url')) if pics else (it.get('thumbnail') or '')
            vars_ = it.get('variations') or []
            for v in (vars_ if vars_ else [None]):
                combo = tuple(a.get('value_name', '') for a in (v.get('attribute_combinations') or [])) if v else ()
                vk = (k, combo)
                if vk in vistas:
                    continue
                vistas.add(vk)
                st = (v.get('available_quantity') if v else it.get('available_quantity')) or 0
                if st <= 0:
                    continue
                m['stock'] += st
                for a in ((v or {}).get('attribute_combinations') or []):
                    val = (a.get('value_name') or '').strip()
                    if not val:
                        continue
                    destino = 'colores' if re.search(r'color', a.get('name', ''), re.I) else 'equipos'
                    if val not in m[destino]:
                        m[destino].append(val)
    return [m for m in modelos.values() if m['stock'] > 0]


def _lista(valores, tope, vacio="—"):
    """Enumera hasta `tope` valores y resume el resto.

    Sin tope hay publicaciones con 20 variantes que desbordan la celda y quedan
    cortadas a la mitad: se ve peor que no listarlas, y la fila deja de ser
    comparable con las demás.
    """
    v = [x for x in valores if x]
    if not v:
        return vacio
    if len(v) <= tope:
        return ", ".join(v)
    return ", ".join(v[:tope]) + f" y {len(v) - tope} más"


def _siguiente_numero(base, serie, arranque=34):
    """Correlativo por año, persistido en disco.

    Arranca en un número intermedio a propósito: un presupuesto «001» avisa que
    es el primero que se emite, y eso debilita la posición al negociar.
    """
    ruta = os.path.join(base, 'presupuestos-mayoristas.json')
    anio = date.today().year
    try:
        est = json.load(open(ruta, encoding='utf-8'))
    except Exception:
        est = {}
    n = int(est.get(str(anio), arranque - 1)) + 1
    est[str(anio)] = n
    try:
        json.dump(est, open(ruta, 'w', encoding='utf-8'), indent=2)
    except Exception as e:
        print(f"  [numeracion] no se pudo guardar el contador: {e}", file=sys.stderr)
    return f"{serie}-{anio}-{n:03d}"


def _hoja_condiciones(wb, args, numero, hoy, descuentos, S):
    """Hoja 2: condiciones comerciales, garantía y cambios por falla.

    En mayorista la pregunta que siempre llega es "¿y si vienen fallados?".
    Tenerlo escrito de antemano evita improvisar cuando ya hay un reclamo.
    """
    ws = wb.create_sheet("Condiciones")
    ws.sheet_view.showGridLines = False
    b_none = Border()
    b_bot = Border(bottom=S['hair'])

    def c(cell, v=None, *, bold=False, size=10, color=None, bg=None,
          align='left', wrap=False, borde=b_none):
        if v is not None:
            cell.value = v
        cell.font = Font(name=S['F'], size=size, bold=bold, color=color or S['TEXT'])
        if bg:
            cell.fill = PatternFill('solid', start_color=bg)
        cell.alignment = Alignment(horizontal=align, vertical='top' if wrap else 'center', wrap_text=wrap)
        cell.border = borde
        return cell

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 86
    ws.column_dimensions['D'].width = 3

    ws.merge_cells('A1:D1')
    c(ws['A1'], "CONDICIONES COMERCIALES", bold=True, size=16, color=S['WHITE'], bg=S['INK'])
    ws.row_dimensions[1].height = 36
    ws.merge_cells('A2:D2')
    c(ws['A2'], f"Presupuesto {numero}   ·   {hoy.strftime('%d/%m/%Y')}   ·   "
                f"{args.empresa} — {args.vendedor}", size=9, color=S['MUTED'])
    ws.row_dimensions[2].height = 20

    esc = " · ".join(f"desde {d} u −{p:g}%" for d, p in descuentos)
    bloques = [
        ("Precios y validez", [
            f"Los precios son mayoristas, en pesos y con IVA incluido, y rigen por {args.validez} días desde la fecha de emisión.",
            "Pasado ese plazo se confirman antes de facturar. Cada modelo tiene su propio precio unitario.",
        ]),
        ("Descuento por volumen", [
            f"Se aplica sobre el total de unidades del pedido, sin importar cómo se repartan entre modelos: {esc}.",
        ]),
        ("Pedido mínimo", [
            f"{args.minimo} unidades por pedido." if args.minimo else "Sin mínimo.",
            "Se pueden combinar todos los modelos y colores de la lista para alcanzarlo.",
        ]),
        ("Forma de pago", [
            f"{args.pago}.",
            "El pedido se considera confirmado una vez acreditado el pago.",
        ]),
        ("Plazo de entrega", [
            f"El pedido se despacha dentro de los {args.entrega} días hábiles posteriores a la acreditación del pago.",
            "El plazo de tránsito del transporte corre por separado y depende del destino.",
        ]),
        ("Envío", [
            f"{args.envio}. El costo del envío no está incluido en los precios de esta lista.",
            "La mercadería viaja por cuenta y riesgo del comprador desde que se entrega al transporte.",
        ]),
        ("Cambios por falla", [
            "Se repone toda unidad con falla de fabricación.",
            "El reclamo se hace dentro de los 7 días corridos de recibida la mercadería, con foto o video de la falla.",
            "La reposición se entrega junto con el siguiente pedido, o se acredita a cuenta si no hubiera uno.",
            "No se cubren daños por uso, colocación forzada, golpes ni desgaste.",
        ]),
        ("Cambios sin falla", [
            "No se aceptan devoluciones ni cambios por diferencia de criterio en color, modelo o cantidad.",
            "Ante la duda de compatibilidad, consultá antes de confirmar el pedido.",
        ]),
        ("Colores y fotos", [
            "Las fotos son ilustrativas. Puede haber variaciones leves de tono entre lotes de fabricación.",
        ]),
        ("Disponibilidad", [
            "La lista se confirma modelo por modelo al momento de tomar el pedido.",
            "Si algún modelo no estuviera disponible se ofrece reemplazo o se descuenta del total.",
        ]),
        ("Otros modelos", [
            "Trabajamos equipos y terminaciones que no figuran en esta lista.",
            "Consultá por el modelo que necesites y te cotizamos.",
        ]),
    ]

    r = 4
    for titulo, lineas in bloques:
        ws.row_dimensions[r].height = 22
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c(ws.cell(row=r, column=2), titulo, bold=True, size=11, color=S['INK'],
          borde=Border(bottom=Side(style='medium', color=S['INK'])))
        r += 1
        for ln in lineas:
            ws.row_dimensions[r].height = 30 if len(ln) > 95 else 18
            c(ws.cell(row=r, column=2), "", borde=b_bot)
            c(ws.cell(row=r, column=3), ln, size=10, color=S['MUTED'], wrap=True, borde=b_bot)
            r += 1
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c(ws.cell(row=r, column=2), "Consultas y pedidos", bold=True, size=11, color=S['INK'],
      borde=Border(bottom=Side(style='medium', color=S['INK'])))
    ws.row_dimensions[r].height = 22
    datos = [x for x in [args.vendedor, args.whatsapp, args.email, args.web] if x]
    r += 1
    ws.row_dimensions[r].height = 24
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c(ws.cell(row=r, column=2), "   ·   ".join(datos), bold=True, size=11,
      color=S['ACCENT'], bg=S['BAND'])

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f'A1:D{r}'
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.oddFooter.left.text = f"{args.empresa} — Presupuesto {numero}"
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--output', required=True)
    p.add_argument('--base', default=os.path.dirname(os.path.abspath(__file__)))
    # El precio mayorista se calcula POR MODELO como un % del precio de venta al
    # público. Un precio único para todo el catálogo rompe por los dos lados:
    # regala los modelos caros (una tarjetera de $55.000 vendida a $26.000) y
    # deja los baratos por ENCIMA del precio público (uno de $16.000 quedaba a
    # $26.000). El % sale de cuánto le deja ML a un revendedor monotributista:
    # por debajo de eso no puede venderlo más barato que vos y ganar plata.
    p.add_argument('--ratio', type=float, default=0.773,
                   help='precio mayorista = precio público × ratio')
    p.add_argument('--desc', nargs='+', default=['50:5', '100:10'],
                   help='desde_unidades:descuento%% por volumen')
    p.add_argument('--incluir-protectores', action='store_true')
    p.add_argument('--validez', type=int, default=15, help='días de validez de la lista')
    p.add_argument('--entrega-dias', dest='entrega', type=int, default=8,
                   help='días hábiles desde acreditado el pago hasta el despacho')
    # Condiciones y contacto: sin esto el presupuesto no se puede aceptar —
    # el cliente no sabe cómo pedir, cuál es el mínimo ni cómo se paga.
    p.add_argument('--empresa',  default='WZMALLAS')
    p.add_argument('--vendedor', default='Rodrigo Zapata')
    p.add_argument('--whatsapp', default='+54 9 2304 216009')
    p.add_argument('--email',    default='znrodrigo23@gmail.com')
    p.add_argument('--web',      default='wzmallas.com')
    p.add_argument('--cliente',  default='', help='nombre del cliente, para encabezar el presupuesto')
    # Numeración correlativa. El contador vive en un archivo para que cada
    # presupuesto salga con su número y no se repita entre generaciones.
    p.add_argument('--numero', default='', help='forzar número (si no, toma el correlativo)')
    p.add_argument('--serie',  default='MAY', help='prefijo de la numeración')
    p.add_argument('--minimo',   type=int, default=20, help='pedido mínimo en unidades (0 = sin mínimo)')
    p.add_argument('--pago',     default='Pago por transferencia, 100% anticipado')
    p.add_argument('--envio',    default='Envío a coordinar')
    args = p.parse_args()

    numero = args.numero or _siguiente_numero(args.base, args.serie)
    descuentos = sorted([(int(d.split(':')[0]), float(d.split(':')[1])) for d in args.desc])
    modelos = [m for m in leer_modelos(args.base, args.incluir_protectores)
               if m['precio_ml'] > 0]
    for m in modelos:
        # Redondeo a $500 para que la lista se lea como una lista de precios
        m['mayorista'] = int(round(m['precio_ml'] * args.ratio / 500) * 500)
    modelos.sort(key=lambda m: -m['stock'])
    if not modelos:
        print('Sin modelos con stock', file=sys.stderr)
        sys.exit(1)

    # ── Sistema visual (mismo que la orden de compra) ────────────
    INK, ACCENT, BAND = "111827", "155E75", "F3F4F6"
    LINE, MUTED, TEXT, WHITE = "E5E7EB", "6B7280", "111827", "FFFFFF"
    F = 'Calibri'
    hair = Side(style='thin', color=LINE)
    b_bot, b_none = Border(bottom=hair), Border()

    def celda(c, v=None, *, bold=False, size=10, color=TEXT, bg=None,
              align='left', fmt=None, wrap=False, borde=b_bot):
        if v is not None:
            c.value = v
        c.font = Font(name=F, size=size, bold=bold, color=color)
        if bg:
            c.fill = PatternFill('solid', start_color=bg)
        c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=wrap)
        c.border = borde
        if fmt:
            c.number_format = fmt
        return c

    wb = Workbook()
    ws = wb.active
    ws.title = "Lista mayorista"
    ws.sheet_view.showGridLines = False

    # Sin columna de stock: la entrega es a plazo fijo para todo el pedido, así
    # que mostrar existencias invita a pedir "eso mandámelo ya" y encima le
    # revela a un posible competidor el tamaño de la operación.
    COLS = [("", 13), ("Modelo", 46), ("Colores disponibles", 32),
            ("Compatible con", 32),
            ("Precio\nde lista", 12), ("Tu precio\ncon volumen", 13),
            ("Cantidad", 11), ("Importe", 15)]
    NC = len(COLS)
    ULT = get_column_letter(NC)

    ws.merge_cells(f'A1:{ULT}1')
    celda(ws['A1'], f"{args.empresa} — LISTA MAYORISTA", bold=True, size=18,
          color=WHITE, bg=INK, align='left', borde=b_none)
    ws.row_dimensions[1].height = 40

    hoy = date.today()
    ws.merge_cells(f'A2:{ULT}2')
    celda(ws['A2'],
          f"Presupuesto {numero}"
          + (f"   ·   Para: {args.cliente}" if args.cliente else "")
          + f"   ·   Emitido {hoy.strftime('%d/%m/%Y')}   ·   Precios en pesos, IVA incluido"
          f"   ·   Validez: {args.validez} días",
          size=9, color=MUTED, bg=WHITE, align='left', borde=b_none)
    ws.row_dimensions[2].height = 20

    # ── Condiciones comerciales ──────────────────────────────────
    # Un presupuesto sin condiciones ni forma de contacto no se puede aceptar:
    # el cliente no sabe cómo pedir, cuánto es el mínimo ni cuándo le llega.
    cond = [
        f"Entrega: {args.entrega} días hábiles desde acreditado el pago",
        f"Pedido mínimo: {args.minimo} unidades" if args.minimo else None,
        args.pago,
        args.envio,
    ]
    ws.merge_cells(f'A3:{ULT}3')
    celda(ws['A3'], "   ·   ".join(c for c in cond if c),
          bold=True, size=10, color=INK, bg="EAF2F6", align='left', borde=b_none)
    ws.row_dimensions[3].height = 24

    contacto = "   ·   ".join(c for c in [
        args.vendedor and f"Contacto: {args.vendedor}",
        args.whatsapp and f"WhatsApp {args.whatsapp}",
        args.email or None,
        args.web or None,
    ] if c)
    ws.merge_cells(f'A4:{ULT}4')
    celda(ws['A4'],
          (contacto or "Contacto: [completar]") +
          "   ·   Para pedir: completá la columna «Cantidad» y devolvenos este archivo.",
          size=9, color=MUTED, bg="EAF2F6", align='left', borde=b_none)
    ws.row_dimensions[4].height = 20

    # ── Bloque de volumen + total en vivo ────────────────────────
    ws.merge_cells('A6:B6'); ws.merge_cells('A7:B7'); ws.merge_cells('A8:B8')
    celda(ws['A6'], "Unidades del pedido", size=10, color=MUTED, bg=BAND, align='left', borde=b_none)
    celda(ws['A7'], "Descuento por volumen", bold=True, size=10, color=INK, bg=BAND, align='left', borde=b_none)
    celda(ws['A8'], "TOTAL DEL PEDIDO", bold=True, size=11, color=INK, bg=BAND, align='left', borde=b_none)

    prim, ultimo = 11, 11 + len(modelos) - 1
    ws['C6'] = f"=SUM(G{prim}:G{ultimo})"
    celda(ws['C6'], size=12, bold=True, color=INK, bg=BAND, align='right', fmt='#,##0', borde=b_none)
    formula = "0"
    for desde, pct in descuentos:
        formula = f"IF($C$6>={desde},{pct / 100},{formula})"
    ws['C7'] = f"={formula}"
    celda(ws['C7'], size=12, bold=True, color=ACCENT, bg=BAND, align='right', fmt='0%', borde=b_none)
    ws['C8'] = f"=SUM(H{prim}:H{ultimo})"
    celda(ws['C8'], size=14, bold=True, color=ACCENT, bg=BAND, align='right', fmt='"$"#,##0', borde=b_none)

    escala = "Cada modelo tiene su propio precio.   ·   " + "   ·   ".join(
        f"desde {d} u: −{pct:g}% en todo el pedido" for d, pct in descuentos)
    ws.merge_cells(f'E6:{ULT}6')
    celda(ws['E6'], "Escala por volumen", bold=True, size=10, color=INK, bg=BAND, align='left', borde=b_none)
    ws.merge_cells(f'E7:{ULT}8')
    celda(ws['E7'], escala + "\nCargá las cantidades y el precio de cada fila se ajusta solo.",
          size=10, color=MUTED, bg=BAND, align='left', wrap=True, borde=b_none)
    for r in (6, 7, 8):
        ws.row_dimensions[r].height = 22
        for col in range(1, NC + 1):
            cc = ws.cell(row=r, column=col)
            if not cc.fill.patternType:
                cc.fill = PatternFill('solid', start_color=BAND)
    ws.row_dimensions[9].height = 8

    HDR = 10
    ws.row_dimensions[HDR].height = 28
    borde_hdr = Border(bottom=Side(style='medium', color=INK))
    for i, (h, w) in enumerate(COLS, 1):
        al = 'right' if i >= 5 else 'left'
        celda(ws.cell(row=HDR, column=i, value=h), bold=True, size=9, color=INK,
              bg=WHITE, align=al, wrap=True, borde=borde_hdr)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f'A{HDR + 1}'

    fotos = {}
    r = HDR + 1
    for m in modelos:
        ws.row_dimensions[r].height = 60
        buf = _miniatura(m['foto'], fotos)
        if buf is not None:
            img = XLImage(buf)
            img.width = img.height = 68
            ws.add_image(img, f'A{r}')
        celda(ws.cell(row=r, column=1), bg=WHITE)
        celda(ws.cell(row=r, column=2, value=m['titulo']), size=10, wrap=True)
        celda(ws.cell(row=r, column=3, value=_lista(m['colores'], 7)),
              size=9, color=MUTED, wrap=True)
        # Se fabrican otros equipos a pedido: decirlo en cada fila abre la puerta
        # a modelos que no están publicados sin tener que listarlos uno por uno.
        compat = _lista(m['equipos'], 7, vacio="")
        compat = (compat + " · Consultar otros modelos") if compat else "Consultar modelos disponibles"
        celda(ws.cell(row=r, column=4, value=compat), size=9, color=MUTED, wrap=True)
        celda(ws.cell(row=r, column=5, value=m['mayorista']), size=10, color=MUTED,
              align='right', fmt='"$"#,##0')
        celda(ws.cell(row=r, column=6, value=f"=ROUND(E{r}*(1-$C$7),-2)"), bold=True,
              size=10, color=ACCENT, align='right', fmt='"$"#,##0')
        celda(ws.cell(row=r, column=7), size=11, bold=True, align='right', fmt='#,##0',
              bg="FFFDF5", borde=Border(bottom=hair, left=hair, right=hair))
        celda(ws.cell(row=r, column=8, value=f"=IF(G{r}=\"\",\"\",G{r}*F{r})"),
              size=10, align='right', fmt='"$"#,##0')
        r += 1

    tr = r
    ws.row_dimensions[tr].height = 26
    borde_tot = Border(top=Side(style='medium', color=INK))
    for col in range(1, NC + 1):
        ws.cell(row=tr, column=col).border = borde_tot
    ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=5)
    celda(ws.cell(row=tr, column=1, value=f"TOTAL — {len(modelos)} modelos"),
          bold=True, size=10, color=INK, align='left', borde=borde_tot)
    celda(ws.cell(row=tr, column=7, value=f"=SUM(G{prim}:G{ultimo})"),
          bold=True, size=11, align='right', fmt='#,##0', borde=borde_tot)
    celda(ws.cell(row=tr, column=8, value=f"=SUM(H{prim}:H{ultimo})"),
          bold=True, size=12, color=ACCENT, align='right', fmt='"$"#,##0', borde=borde_tot)

    nota = tr + 2
    ws.merge_cells(f'A{nota}:{ULT}{nota}')
    celda(ws.cell(row=nota, column=1),
          f"Los precios de esta lista tienen una validez de {args.validez} días y están sujetos a confirmación al momento del pedido. "
          "El descuento por volumen se calcula sobre el total de unidades del pedido, no por modelo. "
          "Trabajamos otros equipos y terminaciones a pedido: consultanos por el modelo que necesites. "
          "Las condiciones completas están en la hoja «Condiciones».",
          size=8, color=MUTED, bg=WHITE, align='left', wrap=True, borde=b_none)
    ws.row_dimensions[nota].height = 26

    # Impresión / PDF: sin esto la tabla (8 columnas) se parte a lo ancho y
    # salen 27 páginas sueltas. Apaisado, ajustado al ancho y repitiendo el
    # encabezado en cada hoja.
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f'{HDR}:{HDR}'
    ws.print_area = f'A1:{ULT}{nota}'
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.left.text = f"{args.empresa} — Presupuesto {numero}"
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8
    ws.oddFooter.left.color = ws.oddFooter.right.color = MUTED

    _hoja_condiciones(wb, args, numero, hoy, descuentos,
                      dict(INK=INK, ACCENT=ACCENT, BAND=BAND, MUTED=MUTED,
                           WHITE=WHITE, TEXT=TEXT, F=F, hair=hair))
    wb.save(args.output)
    print(json.dumps({
        "presupuesto": numero,
        "modelos": len(modelos),
        "unidades_disponibles": sum(m['stock'] for m in modelos),
        "fotos": sum(1 for m in modelos if m['foto']),
        "ratio": args.ratio,
        "descuentos": [{"desde": d, "pct": pct} for d, pct in descuentos],
        "precio_min": min(m['mayorista'] for m in modelos),
        "precio_max": max(m['mayorista'] for m in modelos),
    }, ensure_ascii=False))


if __name__ == '__main__':
    main()
